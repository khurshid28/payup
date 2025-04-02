import os, re

import pythoncom
import qrcode
import win32com.client
from openpyxl import load_workbook
from openpyxl.drawing.image import Image
from django.conf import settings

from stepform.models import GeneratedDocument


class ExcelToPDFConverter:
    def __init__(self, filename, xlsx, global_ip):
        self.filename = filename
        self.global_ip = global_ip
        self.xlsx = xlsx
        self.qr_file_path = ""
        self.pdf_paths = []

    def create_qr_code(self, pdf_filename):
        """ QR kodni yaratadi va saqlaydi """
        qr_file_name = f"{pdf_filename}.png"
        self.qr_file_path = os.path.join(settings.MEDIA_ROOT, "uploads/generated/qrcode", qr_file_name)

        os.makedirs(os.path.dirname(self.qr_file_path), exist_ok=True)

        qr = qrcode.QRCode(
            version=4,
            error_correction=qrcode.constants.ERROR_CORRECT_L,
            box_size=10,
            border=4,
        )

        # _qr va undan keyingi raqamlarni olib tashlash
        clean_name = re.sub(r"_qr\d+$", "", pdf_filename)
        qr.add_data(f"http://{self.global_ip}/media/uploads/generated/pdf/{clean_name}.pdf")
        qr.make(fit=True)

        img = qr.make_image(fill="black", back_color="white")
        img.save(self.qr_file_path)

    def add_qr_to_excel(self, sheet_name, img_positions):
        """Bir nechta QR kodni Excel faylga joylashtiradi"""
        wb = load_workbook(filename=self.xlsx, read_only=False, data_only=True)
        ws = wb[sheet_name]

        for idx, pos in enumerate(img_positions):
            # qr_filename = f"{self.filename}_{sheet_name}_qr{idx + 1}.png"
            qr_filename = f"{self.filename}_{sheet_name}_qr{idx + 1}"
            self.create_qr_code(qr_filename)
            qr_filename_with_png = f"{qr_filename}.png"
            qr_path = os.path.join(settings.MEDIA_ROOT, "uploads/generated/qrcode", qr_filename_with_png)
            img = Image(qr_path)
            img.width = 80
            img.height = 80
            ws.add_image(img, pos)

        excel_file_name = f"{self.filename}_{sheet_name}.xlsx"
        excel_path = os.path.join(settings.MEDIA_ROOT, "uploads/generated/excel", excel_file_name)
        wb.save(excel_path)

        return excel_path

    def convert_excel_to_pdf(self, excel_path, pdf_filename, sheet_name, start_cell, end_cell):
        """ Excel faylni PDF formatiga o‘giradi """
        pythoncom.CoInitialize()
        excel = win32com.client.Dispatch("Excel.Application")
        excel.Visible = False

        wb = excel.Workbooks.Open(excel_path)
        ws = wb.Sheets(sheet_name)

        # Sahifa sozlamalari
        try:
            ws.PageSetup.PaperSize = 9  # xlPaperA4
        except Exception as e:
            print(f"PaperSize sozlashda xatolik: {e}")

        ws.PageSetup.Orientation = 1
        ws.PageSetup.Zoom = False
        ws.PageSetup.FitToPagesWide = 1
        ws.PageSetup.FitToPagesTall = False

        ws.PageSetup.LeftMargin = excel.InchesToPoints(0.3)
        ws.PageSetup.RightMargin = excel.InchesToPoints(0.3)
        ws.PageSetup.TopMargin = excel.InchesToPoints(0.3)
        ws.PageSetup.BottomMargin = excel.InchesToPoints(0.3)

        # ✅ Chop etish maydonini aniq belgilash (xatolik oldini olish)
        try:
            # used_range = ws.UsedRange
            # start_cell = used_range.Cells(1, 1).Address
            # end_cell = used_range.Cells(used_range.Rows.Count, used_range.Columns.Count).Address
            ws.PageSetup.PrintArea = f"{start_cell}:{end_cell}"
            print(f"Chop etish maydoni: {start_cell}:{end_cell}")

        except Exception as e:
            print(f"PrintArea sozlashda xatolik: {e}")
            ws.PageSetup.PrintArea = "A1:G50"  # Zaxira maydon
        # ✅ Ortiqcha varoqlarni olib tashlash
        ws.ResetAllPageBreaks()

        # PDF'ga o‘girish
        pdf_path = os.path.join(settings.MEDIA_ROOT, f"uploads/generated/pdf/{pdf_filename}.pdf")
        ws.ExportAsFixedFormat(0, pdf_path)

        wb.Close(SaveChanges=False)
        excel.Quit()
        del excel
        pythoncom.CoUninitialize()

        return pdf_path

    def generate_multiple_pdfs_with_qr(self):
        """ Bir nechta sahifalar uchun alohida PDFlar yaratadi """
        sheets_info = {
            "shartnoma": {"img_positions": ["A124"], "start_cell": "A1", "end_cell": "B133"},
            "buyruq": {"img_positions": ["C27"], "start_cell": "A1", "end_cell": "E29"},
            "bayonnoma": {"img_positions": ["C40", "C44", "C49"], "start_cell": "A1", "end_cell": "E52"},
            "xulosa": {"img_positions": ["C54", "C58", "C62", "C66"], "start_cell": "A1", "end_cell": "E70"},
            "dalolatnoma": {"img_positions": ["D51"], "start_cell": "A1", "end_cell": "E65"},
            "grafik": {"img_positions": ["A59"], "start_cell": "A1", "end_cell": "G63"},
            "ariza": {"img_positions": ["C20"], "start_cell": "A1", "end_cell": "D25"},
            "muqova": {"start_cell": "A1", "end_cell": "F40"},
            "mijoz_anketasi": {"start_cell": "A1", "end_cell": "F36"},
            "majburiyatnoma": {"start_cell": "A1", "end_cell": "B11"},
            "kredit_ariza": {"img_positions": ["C43"], "start_cell": "A1", "end_cell": "D45"},
            "garov_ariza": {"img_positions": ["C23"], "start_cell": "A1", "end_cell": "C25"},
            "akt_monitoring_1": {"img_positions": ["F21", "F26",], "start_cell": "A1", "end_cell": "H32"},
            "akt_monitoring_2": {"img_positions": ["F21", "F26",], "start_cell": "A1", "end_cell": "H32"},
            "akt_monitoring_3": {"img_positions": ["F21", "F26",], "start_cell": "A1", "end_cell": "H32"},
            "akt_monitoring_4": {"img_positions": ["F21", "F26",], "start_cell": "A1", "end_cell": "H30"},
        }

        self.pdf_paths = {}  # PDF pathlarni dict ko'rinishida saqlaymiz
        for sheet, info in sheets_info.items():
            pdf_filename = f"{self.filename}_{sheet}"

            if "img_positions" in info:
                self.create_qr_code(pdf_filename)
                excel_path = self.add_qr_to_excel(sheet, img_positions=info["img_positions"])
            else:
                excel_path = os.path.join(settings.MEDIA_ROOT, f"uploads/xlsx_template/{self.filename}.xlsx")

            pdf_path = self.convert_excel_to_pdf(
                excel_path, pdf_filename, sheet, info["start_cell"], info["end_cell"]
            )
            self.pdf_paths[sheet] = pdf_path  # Dict shaklida saqlash
            print(sheet)
        print(self.pdf_paths)
        return self.pdf_paths

    def save_to_generated_document(self, application_id, user_id):
        """ Yaratilgan PDF fayllarni GeneratedDocument modeliga saqlaydi """

        def get_relative_path(full_path):
            """ To'liq yo'ldan nisbiy yo‘lni ajratib olish """
            return os.path.relpath(full_path, settings.MEDIA_ROOT)

        document = GeneratedDocument(
            created_by=user_id,
            application_id=application_id,
            pdf_shartnoma=get_relative_path(self.pdf_paths.get('shartnoma')),
            pdf_buyruq=get_relative_path(self.pdf_paths.get('buyruq')),
            pdf_bayonnoma=get_relative_path(self.pdf_paths.get('bayonnoma')),
            pdf_xulosa=get_relative_path(self.pdf_paths.get('xulosa')),
            pdf_dalolatnoma=get_relative_path(self.pdf_paths.get('dalolatnoma')),
            pdf_grafik=get_relative_path(self.pdf_paths.get('grafik')),
            pdf_ariza=get_relative_path(self.pdf_paths.get('ariza')),
            pdf_muqova=get_relative_path(self.pdf_paths.get('muqova')),
            pdf_mijoz_anketasi=get_relative_path(self.pdf_paths.get('mijoz_anketasi')),
            pdf_majburiyatnoma=get_relative_path(self.pdf_paths.get('majburiyatnoma')),
            pdf_kredit_ariza=get_relative_path(self.pdf_paths.get('kredit_ariza')),
            pdf_garov_ariza=get_relative_path(self.pdf_paths.get('garov_ariza')),
            pdf_akt_monitoring_1=get_relative_path(self.pdf_paths.get('akt_monitoring_1')),
            pdf_akt_monitoring_2=get_relative_path(self.pdf_paths.get('akt_monitoring_2')),
            pdf_akt_monitoring_3=get_relative_path(self.pdf_paths.get('akt_monitoring_3')),
            pdf_akt_monitoring_4=get_relative_path(self.pdf_paths.get('akt_monitoring_4')),
        )
        document.save()


    def clear_generated_excel_files(self):
        """ media/uploads/generated/excel papkasidagi barcha fayllarni o'chiradi """
        excel_folder_path = os.path.join(settings.MEDIA_ROOT, "uploads", "generated", "excel")

        if os.path.exists(excel_folder_path):
            for filename in os.listdir(excel_folder_path):
                file_path = os.path.join(excel_folder_path, filename)
                try:
                    if os.path.isfile(file_path):
                        os.remove(file_path)
                        print(f"✅ O'chirildi: {file_path}")
                except Exception as e:
                    print(f"⚠️ Xatolik: {e}")
        else:
            print("📁 Papka topilmadi!")

    def force_kill_excel(self):
        """ Osilib qolgan EXCEL jarayonlarini majburan o'chiradi """
        try:
            os.system("taskkill /f /im excel.exe")
            print("✅ Barcha osilib qolgan Excel jarayonlari tozalandi.")
        except Exception as e:
            print(f"⚠️ Excel jarayonlarini tozalashda xatolik: {e}")